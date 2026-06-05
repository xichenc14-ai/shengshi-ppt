"""
文本框构建器
"""

from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor

from pdf_to_pptx.converter_types import PageContent, TextBlock
from pdf_to_pptx.utils.coord import pdf_bbox_to_pptx
from pdf_to_pptx.utils.color import rgb_to_pptx_rgb

# 字体 fallback 映射
FONT_FALLBACK = {
    "SimHei": "Microsoft YaHei",
    "SimSun": "SimSun",
    "Arial": "Arial",
    "Helvetica": "Calibri",
    "Times": "Times New Roman",
}


class TextBuilder:
    """文本框构建"""

    def __init__(self, slide, page: PageContent):
        self.slide = slide
        self.page = page

    def add_text(self, text_block: TextBlock) -> None:
        """
        将 TextBlock 添加为 PPTX 文本框

        Args:
            text_block: TextBlock 数据
        """
        # 坐标转换
        left, top, width, height = pdf_bbox_to_pptx(
            text_block.bbox,
            self.page.height,
        )

        # 避免负宽高
        if width <= 0 or height <= 0:
            width = Emu(1000000)
            height = Emu(500000)

        # 添加文本框
        txBox = self.slide.shapes.add_textbox(
            left, top, width, height
        )
        tf = txBox.text_frame
        tf.word_wrap = True

        # 设置文本
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = text_block.text

        # 字体
        font_name = self._map_font(text_block.font)
        run.font.name = font_name
        run.font.size = Pt(text_block.size)

        # 颜色
        try:
            run.font.color.rgb = rgb_to_pptx_rgb(text_block.color)
        except Exception:
            pass  # 忽略无效颜色

        # 粗体/斜体
        run.font.bold = text_block.bold
        run.font.italic = text_block.italic

    def _map_font(self, pdf_font: str) -> str:
        """字体名映射"""
        pdf_font_lower = pdf_font.lower()

        for key, mapped in FONT_FALLBACK.items():
            if key.lower() in pdf_font_lower:
                return mapped

        # 无法映射，保留原名
        return pdf_font