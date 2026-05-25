"""
表格构建器
"""

from pptx.util import Pt, Emu

from pdf_to_pptx.converter_types import TableBlock
from pdf_to_pptx.utils.coord import pdf_bbox_to_pptx


class TableBuilder:
    """表格构建"""

    def __init__(self, slide):
        self.slide = slide

    def add_table(self, table_block: TableBlock) -> None:
        """
        将 TableBlock 添加为 PPTX 表格

        Args:
            table_block: TableBlock 数据
        """
        left, top, width, height = pdf_bbox_to_pptx(
            table_block.bbox,
            page_height=0,  # TODO: 传入 page_height
        )

        if width <= 0:
            width = Emu(5000000)
        if height <= 0:
            height = Emu(2000000)

        rows = min(table_block.rows, 50)  # 限制最大行数
        cols = min(table_block.cols, 20)   # 限制最大列数

        table_shape = self.slide.shapes.add_table(
            rows, cols, left, top, width, height
        )
        table = table_shape.table

        # 填充单元格
        for r in range(rows):
            for c in range(cols):
                cell_text = ""
                if r < len(table_block.cells) and c < len(table_block.cells[r]):
                    cell_text = table_block.cells[r][c] or ""

                cell = table.cell(r, c)
                cell.text = cell_text

                # 简化样式：字号 10
                for paragraph in cell.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(10)